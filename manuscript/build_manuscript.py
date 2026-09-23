"""Build the editable Nemesis manuscript from source markdown and numbered assets."""
from pathlib import Path
from datetime import datetime,timezone
import sys,re,json,math,hashlib,argparse,csv,unicodedata
R=Path(__file__).resolve().parent
sys.path.insert(0,str(R/'python_packages'))
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_BREAK,WD_TAB_ALIGNMENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT,WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.mathtext import MathTextParser

OUT=R/'final';QA=R/'qa';EQ=QA/'equations'
WIDTH=6.80
AUTHORS=[('Shafi Bin Sultan','St. Joseph Higher Secondary School','shafibinsultan0207@gmail.com'),('Sabik Bin Sultan','BAF Shaheen College Kurmitola','sabikbinsultan@gmail.com'),('Safwan Sadad','Greenland Residential School','Avoidsafwan@gmail.com')]
COUNTS={'figures':0,'tables':0,'equations':0}
MANIFEST={'figures':[],'tables':[],'equations':[],'warnings':[]}

def read(p):return Path(p).read_text(encoding='utf-8-sig')
def clean(s):
 parts=re.split(r'(https?://\S+|[A-Za-z]:[\\/]\S+)',s)
 s=''.join(part if re.match(r'https?://|[A-Za-z]:[\\/]',part) else re.sub(r'\bNemesis\b','NEMESIS',part) for part in parts)
 return s.replace('\u2011','-').replace('\u2010','-').replace('\u2212','−').strip()
def fig(path,caption,height=3.0):return '\n\n!FIG['+path+'|'+caption+'|'+str(height)+']\n\n'
def renumber_section(s,number):
 lines=[];sub=0
 for line in s.splitlines():
  if line.startswith('# '):line='## '+str(number)+'. '+line[2:]
  elif line.startswith('## '):
   sub+=1;line='### '+str(number)+'.'+str(sub)+'. '+line[3:]
  lines.append(line)
 return '\n'.join(lines)

def full_captions():
 d={};s=read(R/'text/figurelicenseattributions.md')
 for block in re.split(r'\n## ',s)[1:]:
  name=block.splitlines()[0].strip();m=re.search(r'\*\*Full caption:\*\* (.*?)(?=\n\n)',block,re.S)
  if m:
   caption=m.group(1).strip()
   if 'research/licensed_figures/' in block:
    doi=re.search(r'Article and credit link: (\S+)',block)
    if doi:caption+=' '+doi.group(1)+'; https://creativecommons.org/licenses/by/4.0/.'
   d[name]=caption
 return d

def convert_agent(s):
 captions=full_captions()
 s=re.sub(r'\$\$\s*(.*?)\s*\$\$',lambda m:'!EQ['+re.sub(r'\s+',' ',m.group(1).strip())+']',s,flags=re.S)
 p=r'\[\[FIGURE:\s*(.*?)\]\]\s*\n\s*Caption(?: provenance)?:\s*(.*?)(?=\n\n|$)'
 def match(m):
  path=m.group(1).strip();caption=captions.get(Path(path).name,m.group(2).strip())
  return fig(path,caption,3.1 if 'licensed_figures' in path else 3.0)
 s=re.sub(p,match,s,flags=re.S)
 assert '[[FIGURE:' not in s,'Unconverted agent figure'
 return s

def assemble(partial=False):
 files=[R/'text'/f'root_{i:02}.md' for i in range(1,5)]
 missing=[p.name for p in files if not p.exists()]
 if missing and not partial:raise FileNotFoundError('Waiting for final source parts: '+', '.join(missing))
 text='\n\n'.join(read(p) for p in files if p.exists())
 related=renumber_section(read(R/'research/related-work-draft.md'),3)
 # Include only the two available, licensed comparator photos.
 ledger=json.loads(read(R/'research/figure_reuse_ledger.json'))
 for x in ledger:
  if not x.get('local_asset'):continue
  caption=x['suggested_caption']+' Source: https://doi.org/'+x['doi']+'. Licence: '+x['license_url']
  anchor='## 3.' # insert after the corresponding comparative paragraph
  needle='The telescopic-leg robot' if x['id']=='CMP01' else 'Existing integrated robots'
  paragraphs=related.split('\n\n')
  for i,p in enumerate(paragraphs):
   if p.startswith(needle):paragraphs.insert(i+1,fig(x['local_asset'],caption,2.45));break
  related='\n\n'.join(paragraphs)
 agent=convert_agent(read(R/'text/sensing_ml.md'))
 pos=agent.index('# External-data machine-learning')
 sense=renumber_section(agent[:pos],7);ml=renumber_section(agent[pos:],10)
 inserts=[('### 7.2.',fig('diagrams/canva_05.png','Canva-created gas-sensing chain. Sensor identity, heater state, raw-response acquisition and reference calibration precede interpretation. The diagram specifies proposed interfaces; it does not establish selective gas concentrations.',3.0)),('### 7.3.',fig('diagrams/canva_06.png','Canva-created particulate measurement chain. Airflow, sensor checks and environmental covariates support quality-controlled acquisition. Corrections require evaluation against a reference instrument.',3.0)),('### 7.5.',fig('diagrams/canva_07.png','Canva-created ultraviolet, pressure and temperature measurement interfaces. BMP180 does not provide relative humidity. Ultraviolet conversion depends on the actual GUVA board and its calibration.',3.0)+fig('diagrams/canva_08.png','Canva-created water-assessment workflow. Probe conditioning, reference checks, stabilization and sample identity accompany pH and turbidity acquisition. The unidentified rover modules require their own calibration.',3.0))]
 for prefix,addition in inserts:
  sense=re.sub(r'('+re.escape(prefix)+r'[^\n]*\n)',lambda m:m.group(1)+addition,sense,count=1)
 sense=sense.replace('### 7.5.',fig('analysis/figures/A05_sensor_response_lag.png','Analytical first-order sensor response and motion-lag sensitivity. Time constants and rover speeds are declared scenarios. The curves illustrate why acquisition frequency alone cannot remove physical sensor lag; they are not fitted response tests.',3.0)+'### 7.5.',1)
 sense+=fig('analysis/figures/A07_pH_conditioning.png','Analytical pH electrode slope and interface-conditioning sensitivity. The curves follow declared Nernst and circuit assumptions. Actual electrode offset, slope, leakage and temperature compensation require buffer measurements.',3.0)
 ml=re.sub(r'(### 10.2.[^\n]*\n)',lambda m:m.group(1)+fig('diagrams/canva_14.png','Canva-created calibration-analysis pipeline. Cleaning, chronological partitioning, training-only preprocessing, validation selection and frozen testing are distinct stages. External-data results remain separate from future rover calibration.',3.0),ml,count=1)
 refs=json.loads(read(R/'research/references_verified.json'))
 # Place the complete numeric comparison immediately before results interpretation.
 rows=list(csv.DictReader((R/'analysis/ml_outputs/metrics.csv').open(encoding='utf-8-sig')))
 labels={'DummyMean':'Training mean','Ridge':'Ridge*','RandomForest':'Random forest','HistGradientBoosting':'Gradient boosting'}
 tb='!TABLE[Model|Validation RMSE|Test RMSE|Test MAE|Test R²;'+(';'.join('|'.join([labels[x['model']]]+[f"{float(x[k]):.4f}" for k in ['validation_RMSE','test_RMSE','test_MAE','test_R2']]) for x in rows))+']'
 extra=tb+'\n\n*Ridge was selected by validation RMSE before test evaluation. MAE and RMSE are in mg/m³; R² is dimensionless. All test scores use the same 1,102 held-out observations.\n\n'
 ml=re.sub(r'(### 10.4.[^\n]*\n)',lambda m:m.group(1)+'\n'+extra,ml,count=1)
 text=text.replace('@@RELATED_WORK@@',related).replace('@@SENSING@@',sense).replace('@@ML@@',ml)
 # Keep references cited in this manuscript; retain all leads in the supplemental registry.
 selected=[]
 for rr in refs['sources']:
  labels=[rr['author_label']]
  if 'SEN25' in rr['source_ids']:labels+=['U.S. Geological Survey','USGS']
  if 'ADD04' in rr['source_ids']:labels+=['JCGM','Joint Committee for Guides in Metrology']
  if 'M16' in rr['source_ids']:labels+=['NIST']
  found=any(re.search(re.escape(label)+r'\s*(?:\(|,)?\s*'+re.escape(rr['citation_year']),text,re.I) for label in labels)
  if 'M16' in rr['source_ids'] and 'NIST response-robot' in text:found=True
  if found:selected.append(rr)
 selected.sort(key=lambda item: ''.join(c for c in unicodedata.normalize('NFKD',item['citation_text']).casefold() if c.isalnum()))
 reference_text='## References\n\n'+'\n\n'.join(x['citation_text'] for x in selected)
 MANIFEST['reference_keys']=[x['key'] for x in selected]
 MANIFEST['supplemental_reference_count']=len(refs['sources'])
 text=text.replace('@@REFERENCES@@',reference_text)
 text=re.sub(r'(?m)(^#{1,3} [^\n]+)\n(?!\n)',r'\1\n\n',text)
 # Improve page flow by letting closely related explanatory paragraphs precede diagrams.
 blocks=[b.strip() for b in re.split(r'\n\s*\n',text) if b.strip()]
 for k in range(len(blocks)-1,-1,-1):
  if blocks[k].startswith('!FIG[') and 'research/comparison_' not in blocks[k]:
   take=0;advance_words=0
   for following in blocks[k+1:k+6]:
    allowed=(not following.startswith(('#','!','@@','- '))) or following.startswith('!EQ[')
    amount=len(following.split())
    if allowed and advance_words+amount<=320:take+=1;advance_words+=amount
    else:break
   if take:
    graphic=blocks.pop(k);blocks.insert(k+take,graphic)
 text='\n\n'.join(blocks)
 # Use complete native Canva PDF page renders, preserving titles and legends.
 def printed(m):
  name=m.group(1);caption=m.group(2)
  caption=re.sub(r'^Canva-created\s+','',caption)
  caption='Complete Canva-authored diagram with an embedded figure title and arrow legend. '+caption[:1].upper()+caption[1:]
  return '!FIG[diagrams/'+name+'_print.png|'+caption+'|4.5]'
 text=re.sub(r'!FIG\[diagrams/(canva_\d+)\.png\|([^|]+)\|[^\]]+\]',printed,text)
 # Place the census graphic before the detailed screening prose so the page fills naturally.
 blocks=[b.strip() for b in re.split(r'\n\s*\n',text) if b.strip()]
 idx=next((i for i,b in enumerate(blocks) if b.startswith('!FIG[figures/Fig_L1')),None)
 anchor=next((i for i,b in enumerate(blocks) if b.startswith('A documented rule combined')),None)
 if idx is not None and anchor is not None and anchor<idx:
  b=blocks.pop(idx);blocks.insert(anchor,b)
 text='\n\n'.join(blocks)
 # These large architecture diagrams fit better at the start of their own subsection.
 blocks=[b.strip() for b in re.split(r'\n\s*\n',text) if b.strip()]
 for file,heading in [('canva_02_print.png','### 8.1.'),('canva_11_print.png','### 8.3.')]:
  ix=next(i for i,b in enumerate(blocks) if b.startswith('!FIG[') and file in b)
  graphic=blocks.pop(ix);anchor=next(i for i,b in enumerate(blocks) if b.startswith(heading));blocks.insert(anchor+1,graphic)
 # Put the protection-state diagram after its stopping calculation and interpretation.
 ix=next(i for i,b in enumerate(blocks) if b.startswith('!FIG[') and 'canva_12_print.png' in b)
 graphic=blocks.pop(ix);anchor=next(i for i,b in enumerate(blocks) if b.startswith('## 10.'));blocks.insert(anchor,graphic)
 text='\n\n'.join(blocks)
 text=re.sub(r'(!FIG\[figures/prototype_views.png\|[^|]+\|)[^\]]+',r'\g<1>5.5',text)
 # The gripper detail documents the actual illustrated vial placement.
 detail=R/'blender/M03_gripper_detail.png'
 if detail.exists():text=re.sub(r'(!FIG\[blender/M03_mission.png[^\n]+\])',lambda m:m.group(1)+fig('blender/M03_gripper_detail.png','Blender close-up of the sample vial positioned between the gripper tips in the preserved arm pose. The image documents the illustrated contact arrangement; it does not establish grip force or a successful autonomous pickup.',2.8),text,count=1)
 # Final review adjustments keep equation introductions together and improve print legibility.
 blocks=[b.strip() for b in re.split(r'\n\s*\n',text) if b.strip()]
 def move_after_heading(asset,heading):
  ix=next(i for i,b in enumerate(blocks) if b.startswith('!FIG[') and asset in b);graphic=blocks.pop(ix)
  anchor=next(i for i,b in enumerate(blocks) if b.startswith(heading));blocks.insert(anchor+1,graphic)
 move_after_heading('A03_static_stability','### 6.2.')
 # Lag plot fills the prior page; the Delhi plot then has room for larger site labels.
 ix=next(i for i,b in enumerate(blocks) if b.startswith('!FIG[') and 'A05_sensor_response_lag' in b);graphic=blocks.pop(ix)
 anchor=next(i for i,b in enumerate(blocks) if b.startswith('!FIG[') and 'zheng2019_fig4' in b);blocks.insert(anchor,graphic)
 # Uncertainty formula and definition immediately follow their existing lead-in.
 ix=next(i for i,b in enumerate(blocks) if b.startswith('!EQ[u_c^2'))
 eq=blocks.pop(ix);definition=blocks.pop(ix)
 anchor=next(i for i,b in enumerate(blocks) if b.startswith('For a calibrated measurand'))
 blocks[anchor+1:anchor+1]=[eq,definition]
 # Introduce networking before its large block diagram to fill the table page.
 ix=next(i for i,b in enumerate(blocks) if b.startswith('A local access point allows'))
 paragraph=blocks.pop(ix);anchor=next(i for i,b in enumerate(blocks) if b.startswith('!FIG[') and 'canva_11_print' in b);blocks.insert(anchor,paragraph)
 # Balance large illustrations with their own prose; no block crosses a subsection.
 def move_graphic_relative(asset,paragraph_start,after=False):
  ix=next(i for i,b in enumerate(blocks) if b.startswith('!FIG[') and asset in b);graphic=blocks.pop(ix)
  anchor=next(i for i,b in enumerate(blocks) if b.startswith(paragraph_start));blocks.insert(anchor+int(after),graphic)
 def move_paragraph_after_graphic(paragraph_start,asset):
  ix=next(i for i,b in enumerate(blocks) if b.startswith(paragraph_start));paragraph=blocks.pop(ix)
  anchor=next(i for i,b in enumerate(blocks) if b.startswith('!FIG[') and asset in b);blocks.insert(anchor+1,paragraph)
 move_graphic_relative('prototype_views.png','The reconstruction preserves',True)
 move_graphic_relative('A3P5_Nemesis_Wiring_Detail.png','The visual refinement introduces',True)
 move_graphic_relative('A01_steering_kinematics','Here δ',True)
 move_after_heading('canva_01_print.png','### 4.3.')
 move_graphic_relative('A05_sensor_response_lag','Zheng et al. (2019)')
 move_graphic_relative('canva_07_print.png','For a calibrated measurand')
 move_after_heading('canva_13_print.png','### 9.2.')
 move_graphic_relative('A06_stopping_clearance','A software stop and emergency energy interruption',True)
 ix=next(i for i,b in enumerate(blocks) if b.startswith('!FIG[') and 'canva_12_print' in b);graphic=blocks.pop(ix)
 anchor=next(i for i,b in enumerate(blocks) if b.startswith('!FIG[') and 'A06_stopping_clearance' in b);blocks.insert(anchor+1,graphic)
 move_paragraph_after_graphic('Each task needs a measurable completion condition','M01_mission')
 move_paragraph_after_graphic('The chronological trace and hourly agreement plot','03_selected_model_predictions')
 move_paragraph_after_graphic('Eligible observations covered','04_residual_humidity')
 ix=next(i for i,b in enumerate(blocks) if b.startswith('Eligible observations covered'));paragraph=blocks.pop(ix)
 anchor=next(i for i,b in enumerate(blocks) if b.startswith('The chronological trace and hourly agreement plot'));blocks.insert(anchor+1,paragraph)
 move_graphic_relative('05_desktop_cost_and_sensitivity','Desktop measurements included',True)
 text='\n\n'.join(blocks)
 def resize(m):
  path,caption,old=m.groups();height=float(old)
  if any(name in path for name in ['01_dataset_audit','03_selected_model_predictions','04_residual_humidity']):height=4.8
  elif 'zheng2019_fig4' in path:height=5.2
  elif 'barkjohn2021_fig4' in path:height=4.4
  elif 'M02_mission' in path:height=3.0
  elif 'M02_probe_detail' in path:height=2.4
  elif 'A05_sensor_response_lag' in path or 'M01_mission' in path:
   with Image.open(R/path) as im:w,h=im.size
   height=(6.2 if 'A05_' in path else 3.8)*h/w
  return '!FIG['+path+'|'+caption+'|'+str(round(height,4))+']'
 text=re.sub(r'!FIG\[([^|]+)\|([^|]+)\|([^\]]+)\]',resize,text)
 # Content-independent captions and within-subsection layout controls.
 caption_file=R/'text/caption_overrides.json'
 captions=json.loads(read(caption_file)) if caption_file.exists() else {}
 layout_file=R/'text/layout_overrides.json'
 layout=json.loads(read(layout_file)) if layout_file.exists() else {}
 if 'captions' in captions:captions=captions['captions']
 def final_figure(m):
  path,caption,height=m.groups()
  caption=captions.get(path,captions.get(Path(path).name,caption))
  height=layout.get('figure_heights_in',{}).get(path,layout.get('figure_heights_in',{}).get(Path(path).name,height))
  assert isinstance(caption,str) and '|' not in caption and '\n' not in caption,(path,caption)
  return '!FIG['+path+'|'+caption+'|'+str(height)+']'
 text=re.sub(r'!FIG\[([^|]+)\|([^|]+)\|([^\]]+)\]',final_figure,text)
 blocks=[b.strip() for b in re.split(r'\n\s*\n',text) if b.strip()]
 def subsection_of(index):
  for b in reversed(blocks[:index+1]):
   m=re.match(r'^#{2,3} (\d+(?:\.\d+)*)\.?\s',b)
   if m:return m.group(1)
  return None
 for rule in layout.get('placements',[]):
  asset=rule['asset'];ix=next(i for i,b in enumerate(blocks) if b.startswith('!FIG[') and asset in b)
  origin=subsection_of(ix);graphic=blocks.pop(ix)
  if 'before_paragraph' in rule or 'after_paragraph' in rule:
   prefix=rule.get('before_paragraph',rule.get('after_paragraph'))
   anchor=next(i for i,b in enumerate(blocks) if subsection_of(i)==origin and not b.startswith(('#','!')) and b.startswith(prefix))
   dest=anchor+int('after_paragraph' in rule)
  elif 'before_asset' in rule or 'after_asset' in rule:
   target=rule.get('before_asset',rule.get('after_asset'))
   anchor=next(i for i,b in enumerate(blocks) if b.startswith('!FIG[') and target in b)
   dest=anchor+int('after_asset' in rule)
  elif 'before_table' in rule or 'after_table' in rule:
   prefix=rule.get('before_table',rule.get('after_table'))
   anchor=next(i for i,b in enumerate(blocks) if b.startswith('!TABLE['+prefix))
   dest=anchor+int('after_table' in rule)
  elif rule.get('after_heading'):
   anchor=next(i for i,b in enumerate(blocks) if re.match(r'^#{2,3} '+re.escape(origin)+r'\.\s',b));dest=anchor+1
  elif 'up_content_blocks' in rule:
   dest=ix-int(rule['up_content_blocks'])
   assert dest>=0 and all(not b.startswith(('#','!FIG[','!TABLE[')) for b in blocks[dest:ix]),rule
   # Keep an equation with the paragraph that introduces it.
   if blocks[dest].startswith('!EQ['):dest-=1
  elif 'down_content_blocks' in rule:
   dest=ix+int(rule['down_content_blocks'])
   assert all(not b.startswith(('#','!FIG[','!TABLE[')) for b in blocks[ix:dest]),rule
   if dest<len(blocks) and blocks[dest].startswith('!EQ['):dest+=1
  else:raise ValueError('Unknown placement rule: '+str(rule))
  assert subsection_of(max(0,dest-1))==origin,('Placement crosses subsection',rule,origin,subsection_of(max(0,dest-1)))
  blocks.insert(dest,graphic)
 for rule in layout.get('paragraph_placements',[]):
  ix=next(i for i,b in enumerate(blocks) if not b.startswith(('#','!')) and b.startswith(rule['paragraph']))
  origin=subsection_of(ix);paragraph=blocks.pop(ix)
  if 'before_table' in rule:dest=next(i for i,b in enumerate(blocks) if b.startswith('!TABLE['+rule['before_table']))
  else:raise ValueError('Unknown paragraph placement: '+str(rule))
  assert subsection_of(dest)==origin,('Paragraph crosses subsection',rule)
  blocks.insert(dest,paragraph)
 text='\n\n'.join(blocks)
 MANIFEST['layout_overrides']=layout
 MANIFEST['caption_override_count']=len(captions)
 # Every figure must retain its source-assigned semantic subsection.
 expected={
 'A3P5_Nemesis_Hero.png':'1','Fig_L1_metadata_census.png':'2.2','Fig_L2_metadata_abstract_coverage.png':'2.2',
 'comparison_ResQbot_Saputra2021_Fig10a_CC-BY-4.jpg':'3.3','comparison_telescopic_rover_Mohamed2025_Fig1c_CC-BY-4.jpg':'3.3',
 'prototype_views.png':'4.1','A3P5_Nemesis_Wiring_Detail.png':'4.2','canva_01_print.png':'4.3',
 'canva_03_print.png':'5.1','A01_steering_kinematics.png':'5.1','canva_04_print.png':'5.2','A02_grade_traction_torque.png':'5.3',
 'canva_09_print.png':'6.1','M03_mission.png':'6.1','M03_gripper_detail.png':'6.1','A03_static_stability.png':'6.2','M02_mission.png':'6.3','M02_probe_detail.png':'6.3',
 'canva_05_print.png':'7.2','canva_06_print.png':'7.3','jayaratne2018_fig2.png':'7.3','patel2024_fig4.png':'7.3','A05_sensor_response_lag.png':'7.4','zheng2019_fig4.png':'7.4',
 'canva_07_print.png':'7.5','canva_08_print.png':'7.5','A07_pH_conditioning.png':'7.5','canva_02_print.png':'8.1','A04_energy_endurance.png':'8.1','canva_11_print.png':'8.3',
 'canva_10_print.png':'9.1','M05_mission.png':'9.1','canva_13_print.png':'9.2','M01_mission.png':'9.2','M04_mission.png':'9.2','canva_12_print.png':'9.3','A06_stopping_clearance.png':'9.3',
 '01_dataset_audit.png':'10.1','canva_14_print.png':'10.2','barkjohn2021_fig4.png':'10.3','02_test_metrics.png':'10.4','03_selected_model_predictions.png':'10.4','04_residual_humidity.png':'10.4','05_desktop_cost_and_sensitivity.png':'10.5','A08_parameter_sensitivity.png':'11.1'}
 membership=[];current=None
 for line in text.splitlines():
  m=re.match(r'^#{2,3} (\d+(?:\.\d+)*)\.?\s',line)
  if m:current=m.group(1)
  m=re.match(r'^!FIG\[([^|]+)\|',line)
  if m:
   name=Path(m.group(1)).name;wanted=expected[name]
   membership.append({'figure_number':len(membership)+1,'file':m.group(1),'actual_subsection':current,'expected_subsection':wanted,'matches':current==wanted})
 assert len(membership)==45 and all(x['matches'] for x in membership),membership
 MANIFEST['figure_sections']=membership
 # Resolve all internal callouts only after final figure ordering is known.
 callouts=json.loads(read(R/'text/figure_callouts.json'))['callouts']
 byfile={x['file']:x for x in membership};seen=[];callout_audit=[]
 blocks=[b.strip() for b in re.split(r'\n\s*\n',text) if b.strip()]
 for item in callouts:
  entries=[byfile[name] for name in item['figure_files']]
  assert all(x['actual_subsection']==item['subsection'] for x in entries),item
  nums=sorted(x['figure_number'] for x in entries)
  if len(nums)==1:label='Figure '+str(nums[0])
  elif len(nums)>2 and nums==list(range(nums[0],nums[-1]+1)):label='Figures '+str(nums[0])+'–'+str(nums[-1])
  elif len(nums)==2:label='Figures '+str(nums[0])+' and '+str(nums[1])
  else:label='Figures '+', '.join(map(str,nums[:-1]))+' and '+str(nums[-1])
  sentence=item['template'].replace('{figures}',label)
  current=None;target=None
  for i,block in enumerate(blocks):
   m=re.match(r'^#{2,3} (\d+(?:\.\d+)*)\.?\s',block)
   if m:current=m.group(1)
   if current!=item['subsection'] or block.startswith(('#','!','*','- ')):continue
   following=blocks[i+1] if i+1<len(blocks) else ''
   if following.startswith('!EQ[') or block.endswith(':'):continue
   target=i;break
  assert target is not None,('No suitable callout paragraph',item)
  blocks[target]+=' '+sentence
  seen.extend(x['file'] for x in entries)
  callout_audit.append({'subsection':item['subsection'],'figure_numbers':nums,'figure_files':item['figure_files'],'sentence':sentence})
 assert len(seen)==45 and len(set(seen))==45 and set(seen)==set(byfile)
 MANIFEST['figure_callouts']={'all_figures_covered':True,'figure_count':len(seen),'callout_count':len(callout_audit),'callouts':callout_audit}
 text='\n\n'.join(blocks)
 text=re.sub(r'(?<!\w)N m(?!\w)','N·m',text)
 if not partial:assert '@@' not in text,'Unresolved manuscript insert'
 return text

def font(style,size,bold=False,italic=False):
 style.font.name='Times New Roman';style.font.size=Pt(size);style.font.bold=bold;style.font.italic=italic;style.font.color.rgb=RGBColor(0,0,0)
 rf=style.element.get_or_add_rPr().rFonts
 for key in list(rf.attrib):
  if 'theme' in key.lower():del rf.attrib[key]
 for key in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+key),'Times New Roman')
 for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)

def setup():
 d=Document();sec=d.sections[0]
 sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
 sec.left_margin=sec.right_margin=Inches((8.27-WIDTH)/2)
 sec.top_margin=Inches(.72);sec.bottom_margin=Inches(.68);sec.header_distance=Inches(.30);sec.footer_distance=Inches(.30)
 for name,size,bold,italic in [('Normal',11,False,False),('Title',20.5,True,False),('Heading 1',12,True,False),('Heading 2',11,True,False),('Heading 3',11,True,True),('Caption',9.5,False,False),('Header',9,False,True),('Footer',9,False,False)]:font(d.styles[name],size,bold,italic)
 normal=d.styles['Normal'].paragraph_format;normal.line_spacing=1.05;normal.space_after=Pt(5);normal.widow_control=True
 for name in ['Heading 1','Heading 2','Heading 3']:
  p=d.styles[name].paragraph_format;p.space_before=Pt(9 if name=='Heading 1' else 7);p.space_after=Pt(4);p.keep_with_next=True;p.keep_together=True
 d.styles['Title'].paragraph_format.space_after=Pt(8)
 d.styles['Title'].paragraph_format.line_spacing=Pt(25.3)
 d.styles['Caption'].paragraph_format.space_after=Pt(7);d.styles['Caption'].paragraph_format.line_spacing=1.0
 # Explicitly clear every header variant, including first and even pages.
 for section in d.sections:
  for header in [section.header,section.first_page_header,section.even_page_header]:
   for element in list(header._element):header._element.remove(element)
   header._element.append(OxmlElement('w:p'))
 p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='1';r.append(t);f.append(r);p._p.append(f)
 d.core_properties.title=read(R/'text/manuscript_ieee.md').splitlines()[0].lstrip('# ') if (R/'text/manuscript_ieee.md').exists() else 'A3P5 NEMESIS design and analytical assessment of a modular rover for environmental inspection'
 d.core_properties.author='; '.join(a[0] for a in AUTHORS)
 d.core_properties.subject='Engineering design, analytical assessment and external-data calibration benchmark'
 d.core_properties.keywords='environmental robotics; four-wheel steering; calibration; NEMESIS'
 d.core_properties.comments='Photo-constrained rover reconstruction, subsystem design, conditional engineering calculations and an external-data calibration benchmark.'
 d.core_properties.created=datetime(2026,9,15,tzinfo=timezone.utc)
 d.core_properties.modified=datetime.now(timezone.utc)
 return d

def inline(p,text):
 text=clean(text)
 variable=re.compile(r'(?<![\w:.\-])(v_y|R_s|R_L|V_c|V_L|R_0|c_raw|d_lag|E_7|u_model|W_req|M_front|y_i|z_i|x_i|c_dry|Σ_x|V_ref|m_b|m_p|p_b|p_p|T_j)(?![\w])')
 def add(value,bold=False,italic=False):
  for part in variable.split(value):
   if not part:continue
   if variable.fullmatch(part):
    base,sub=part.split('_',1);r=p.add_run(base);r.italic=part not in ['p_b','p_p'];r.bold=bold or part in ['p_b','p_p']
    r=p.add_run(sub);r.font.subscript=True;r.italic=len(sub)==1 and sub.isalpha();r.bold=bold
   else:
    r=p.add_run(part);r.bold=bold;r.italic=italic
    if any(ord(c)>0x1D000 for c in part):r.font.name='Cambria Math'
 for token in re.split(r'(\*\*.*?\*\*|`[^`]+`|\[[^\]]+\]\([^)]*\))',text):
  if not token:continue
  if token.startswith('**') and token.endswith('**'):add(token[2:-2],bold=True)
  elif token.startswith('`') and token.endswith('`'):add(token[1:-1],italic=True)
  elif re.match(r'\[[^\]]+\]\(',token):
   m=re.match(r'\[([^\]]+)\]\((.*?)\)',token);add(m.group(1)+' ('+m.group(2)+')')
  else:add(token)
 return p

def authors(d):
 p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(5);p.paragraph_format.keep_with_next=True
 for i,(name,aff,em) in enumerate(AUTHORS,1):
  if i>1:p.add_run('   ')
  p.add_run(name).bold=True;r=p.add_run(str(i));r.font.superscript=True
 for i,(name,aff,em) in enumerate(AUTHORS,1):
  p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(1);p.paragraph_format.keep_with_next=True
  r=p.add_run(str(i));r.font.superscript=True;r.font.size=Pt(9.5)
  r=p.add_run(' '+aff);r.font.size=Pt(9.5)
 p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(8);p.paragraph_format.keep_with_next=True
 for i,(name,aff,em) in enumerate(AUTHORS,1):
  if i>1:p.add_run('   ')
  r=p.add_run(str(i));r.font.superscript=True;r.font.size=Pt(9)
  r=p.add_run(' '+em);r.font.size=Pt(9)

def addfig(d,s):
 path,caption,maxh=s.split('|',2);path=(R/path).resolve();maxh=float(maxh)
 if not path.exists():raise FileNotFoundError('Missing manuscript figure: '+str(path))
 with Image.open(path) as im:w,h=im.size
 fw=min(WIDTH,maxh*w/h);fh=fw*h/w
 COUNTS['figures']+=1;n=COUNTS['figures']
 p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(3);p.paragraph_format.keep_with_next=True;p.paragraph_format.keep_together=True
 shape=p.add_run().add_picture(str(path),width=Inches(fw));shape._inline.docPr.set('descr',clean(caption));shape._inline.docPr.set('title','Figure '+str(n))
 p=d.add_paragraph(style='Caption');p.paragraph_format.keep_together=True;inline(p,'Figure '+str(n)+'. '+caption)
 MANIFEST['figures'].append({'number':n,'path':str(path),'width_in':round(fw,3),'height_in':round(fh,3),'caption':caption})

def normmath(s):
 s=re.sub(r'\\(mathsf|mathcal|mathbf|boldsymbol|widehat|overline)\s*(\\[A-Za-z]+|[A-Za-z])',lambda m:'\\'+m.group(1)+'{'+m.group(2)+'}',s)
 s=s.replace(r'\frac1n',r'\frac{1}{n}').replace(r'\lVert',r'\Vert').replace(r'\rVert',r'\Vert')
 return s

def addeq(d,source):
 if d.paragraphs:
  previous=d.paragraphs[-1]
  if not previous._p.xpath('.//w:drawing') and previous.style.name!='Caption':previous.paragraph_format.keep_with_next=True
 source=clean(source);s=normmath(source);COUNTS['equations']+=1;n=COUNTS['equations'];p=EQ/f'eq_{n:03}.png'
 cache=EQ/f'eq_{n:03}.sha256';key=hashlib.sha256((s+'|STIX12|360dpi').encode()).hexdigest()
 if not p.exists() or not cache.exists() or read(cache)!=key:
  plt.rcParams.update({'font.family':'serif','mathtext.fontset':'stix'})
  MathTextParser('agg').parse('$'+s+'$',dpi=140)
  f=plt.figure(figsize=(10,.7));f.patch.set_alpha(0);f.text(0,.5,'$'+s+'$',fontsize=12,ha='left',va='center');f.savefig(p,dpi=360,bbox_inches='tight',pad_inches=.035,transparent=True);plt.close(f)
  cache.write_text(key,encoding='utf-8')
 with Image.open(p) as im:w,h=im.size
 fw=min(6.20,w/360);fh=fw*h/w
 para=d.add_paragraph();para.paragraph_format.space_before=Pt(4);para.paragraph_format.space_after=Pt(6);para.paragraph_format.keep_together=True
 para.paragraph_format.tab_stops.add_tab_stop(Inches(WIDTH/2),WD_TAB_ALIGNMENT.CENTER)
 para.paragraph_format.tab_stops.add_tab_stop(Inches(WIDTH),WD_TAB_ALIGNMENT.RIGHT)
 para.add_run('\t');pic=para.add_run().add_picture(str(p),width=Inches(fw));pic._inline.docPr.set('descr','Equation '+str(n)+': '+source);para.add_run('\t('+str(n)+')')
 MANIFEST['equations'].append({'number':n,'latex':source,'normalized_latex':s,'width_in':round(fw,3),'height_in':round(fh,3),'path':str(p)})

def addtable(d,s,heading):
 rows=[[clean(c) for c in row.split('|')] for row in s.split(';')];nc=len(rows[0]);assert all(len(r)==nc for r in rows),rows
 COUNTS['tables']+=1;n=COUNTS['tables']
 title={'Evidence class':'Evidence sources and permitted interpretations','Geometric quantity':'Geometric quantities and analytical assumptions','Interface':'Subsystem interfaces and verification requirements','Model':'Validation and frozen-test comparison on external UCI data'}.get(rows[0][0],re.sub(r'^\d+(?:\.\d+)*[. ]*','',heading))
 p=d.add_paragraph(style='Caption');p.paragraph_format.keep_with_next=True;inline(p,'Table '+str(n)+'. '+title+'.')
 table=d.add_table(rows=1,cols=nc);table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
 weights=[max(8,min(58,sum(len(row[c]) for row in rows)/len(rows))) for c in range(nc)]
 if nc==3:weights=[.25,.29,.46]
 if nc==5:weights=[.27,.20,.18,.18,.17]
 widths=[WIDTH*w/sum(weights) for w in weights]
 for c,width in zip(table.columns,widths):c.width=Inches(width)
 borders=OxmlElement('w:tblBorders')
 for side in ['top','left','bottom','right','insideH','insideV']:
  b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
 table._tbl.tblPr.append(borders)
 margins=OxmlElement('w:tblCellMar')
 for side,val in [('top',65),('left',90),('bottom',65),('right',90)]:
  el=OxmlElement('w:'+side);el.set(qn('w:w'),str(val));el.set(qn('w:type'),'dxa');margins.append(el)
 table._tbl.tblPr.append(margins)
 for ri,row in enumerate(rows):
  cells=table.rows[0].cells if ri==0 else table.add_row().cells
  trpr=cells[0]._tc.getparent().get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
  if ri==0:repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
  for ci,(cell,text) in enumerate(zip(cells,row)):
   cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.0;p.paragraph_format.keep_together=True
   p.alignment=WD_ALIGN_PARAGRAPH.LEFT if len(text)>16 else WD_ALIGN_PARAGRAPH.CENTER
   inline(p,text)
   for run in p.runs:run.font.size=Pt(9.5);run.bold=ri==0
   if ri==0:
    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7E6E6');cell._tc.get_or_add_tcPr().append(shade)
 p=d.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=Pt(3);p.add_run().font.size=Pt(3)
 MANIFEST['tables'].append({'number':n,'caption':title,'rows':len(rows)-1,'columns':nc})

def addcode(d,s):
 path,caption=s.split('|',1)
 p=d.add_paragraph(style='Caption');p.paragraph_format.keep_with_next=True
 inline(p,'Listing 1. '+caption)
 p=d.add_paragraph();p.paragraph_format.keep_together=True
 p.paragraph_format.space_before=Pt(2);p.paragraph_format.space_after=Pt(7)
 p.paragraph_format.line_spacing=Pt(10.5)
 r=p.add_run(read(R/path).rstrip());r.font.name='Courier New';r.font.size=Pt(8.5)
 MANIFEST['code_listing']={'path':path,'caption':caption,'lines':len(read(R/path).splitlines())}

def build(text):
 d=setup();lines=text.splitlines();i=0;buf=[];heading='';references=False;title=True
 def flush():
  nonlocal buf
  if not buf:return
  value=' '.join(buf).strip();buf=[]
  if not value:return
  p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
  if references:
   p.alignment=WD_ALIGN_PARAGRAPH.LEFT
   p.paragraph_format.left_indent=Inches(.22);p.paragraph_format.first_line_indent=Inches(-.22);p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=1.0
  if value.startswith('Keywords:'):
   p.paragraph_format.space_after=Pt(8)
  inline(p,value)
  if value.startswith(('For n held-out observations','For a calibrated measurand','A useful family of humidity corrections','For a conventional divider')) or value.endswith(('using','written','solves')):p.paragraph_format.keep_with_next=True
  if references:
   for r in p.runs:r.font.size=Pt(9)
 while i<len(lines):
  line=lines[i].strip();i+=1
  if not line:flush();continue
  if line.startswith('#'):
   flush();level=len(line)-len(line.lstrip('#'));h=line[level:].strip()
   if title:
    p=d.add_paragraph(style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True;inline(p,h)
    for run in p.runs:run.bold=True
    authors(d);title=False;continue
   heading=h;references=h.lower()=='references'
   h=re.sub(r'^(\d+(?:\.\d+)*)\.\s+',r'\1 ',h)
   p=d.add_paragraph(style='Heading '+str(min(level-1,3)));inline(p,h)
   continue
  if line.startswith('!FIG[') and line.endswith(']'):flush();addfig(d,line[5:-1]);continue
  if line.startswith('!CODE[') and line.endswith(']'):flush();addcode(d,line[6:-1]);continue
  if line.startswith('!EQ[') and line.endswith(']'):flush();addeq(d,line[4:-1]);continue
  if line.startswith('!TABLE[') and line.endswith(']'):flush();addtable(d,line[7:-1],heading);continue
  if line.startswith('- '):flush();p=d.add_paragraph(style='List Bullet');inline(p,line[2:]);continue
  if line.startswith('@@'):raise ValueError('Unresolved marker '+line)
  buf.append(line)
 flush();return d

def main():
 ap=argparse.ArgumentParser();ap.add_argument('--check',action='store_true');ap.add_argument('--partial',action='store_true');a=ap.parse_args()
 revision=R/'text/manuscript_ieee.md'
 text=read(revision) if revision.exists() else assemble(a.partial)
 if revision.exists():
  prior=json.loads(read(QA/'build_manifest.json')) if (QA/'build_manifest.json').exists() else {}
  for field in ['figure_callouts','supplemental_reference_count']:MANIFEST[field]=prior.get(field)
  callout_specs=json.loads(read(R/'text/figure_callouts.json'))['callouts']
  old_sections={Path(asset).name:item['subsection'] for item in callout_specs for asset in item['figure_files']}
  section=None;section_rows=[]
  for line in text.splitlines():
   match=re.match(r'^#{2,3} (\d+(?:\.\d+)*)\.?\s',line)
   if match:section=match.group(1)
   match=re.match(r'^!FIG\[([^|]+)\|',line)
   if match:
    asset=match.group(1);expected=old_sections.get(Path(asset).name,section)
    section_rows.append({'figure_number':len(section_rows)+1,'file':asset,'actual_subsection':section,'expected_subsection':expected,'matches':section==expected})
  assert all(x['matches'] for x in section_rows),'Figure changed subsection'
  MANIFEST['figure_sections']=section_rows
  byfile={x['file']:x['figure_number'] for x in section_rows};callout_rows=[]
  for item in callout_specs:
   nums=sorted(byfile[x] for x in item['figure_files'])
   if len(nums)==1:label='Figure '+str(nums[0])
   elif len(nums)==2:label='Figures '+str(nums[0])+' and '+str(nums[1])
   else:label='Figures '+str(nums[0])+'–'+str(nums[-1])
   sentence=item['template'].replace('{figures}',label)
   if item['subsection']=='7.5':
    sentence='Figure 25 presents the calculated pH-conditioning response, Figure 26 the water-assessment workflow, and Figure 27 the ultraviolet, pressure and temperature interfaces.'
   callout_rows.append({'subsection':item['subsection'],'figure_numbers':nums,'figure_files':item['figure_files'],'sentence':sentence})
  MANIFEST['figure_callouts']={'all_figures_covered':len(byfile)==45,'figure_count':len(byfile),'callout_count':len(callout_rows),'callouts':callout_rows}
  registry=json.loads(read(R/'research/references_verified.json'))
  MANIFEST['supplemental_reference_count']=len(registry) if isinstance(registry,list) else len(registry['sources'])
  MANIFEST['citation_style']='IEEE numeric'
  MANIFEST['reference_keys']=[x['key'] for x in json.loads(read(R/'research/ieee_references.json'))['references']]
  layout_path=R/'text/ieee_layout.json'
  if layout_path.exists():
   sizes=json.loads(read(layout_path)).get('figure_heights_in',{})
   text=re.sub(r'!FIG\[([^|]+)\|([^|]+)\|([^\]]+)\]',lambda m:'!FIG['+m[1]+'|'+m[2]+'|'+str(sizes.get(Path(m[1]).name,m[3]))+']',text)
 if a.check:
  paths=[(R/m.group(1)).resolve() for m in re.finditer(r'^!FIG\[([^|]+)\|',text,re.M)]
  missing=[str(p) for p in paths if not p.exists()]
  eqs=re.findall(r'^!EQ\[(.*)\]$',text,re.M);errors=[]
  for k,s in enumerate(eqs,1):
   try:MathTextParser('agg').parse('$'+normmath(s)+'$',dpi=120)
   except Exception as e:errors.append({'number':k,'latex':s,'error':str(e)})
  print(json.dumps({'source_characters':len(text),'figure_count':len(paths),'missing_figures':missing,'equation_count':len(eqs),'equation_errors':errors},ensure_ascii=False,indent=2));return
 OUT.mkdir(exist_ok=True);QA.mkdir(exist_ok=True);EQ.mkdir(exist_ok=True)
 d=build(text);name='A3P5_NEMESIS_Research_Manuscript.docx'
 if revision.exists():
  sys.path.insert(0,str(R/'research'))
  from style_ieee_docx import style_references
  style_references(d,R/'research/ieee_references.json')
  in_references=False
  for paragraph in d.paragraphs:
   if paragraph.text.strip()=='References':in_references=True;continue
   if not in_references:continue
   for run in paragraph._p.iter(qn('w:r')):
    rpr=run.find(qn('w:rPr'))
    if rpr is None:rpr=OxmlElement('w:rPr');run.insert(0,rpr)
    for tag in ['w:sz','w:szCs']:
     size=rpr.find(qn(tag))
     if size is None:size=OxmlElement(tag);rpr.append(size)
     size.set(qn('w:val'),'20')
 d.save(OUT/name)
 (OUT/'A3P5_NEMESIS_Research_Manuscript.md').write_text(text,encoding='utf-8')
 MANIFEST['counts']=COUNTS.copy();MANIFEST['document']=str(OUT/name);MANIFEST['source_sha256']=hashlib.sha256(text.encode()).hexdigest();MANIFEST['word_count_approx']=len(re.findall(r'\b\w+\b',text));MANIFEST['equation_representation']='High-resolution equation images with source LaTeX and alt text; body text, headings and tables are editable.'
 (QA/'build_manifest.json').write_text(json.dumps(MANIFEST,ensure_ascii=False,indent=2),encoding='utf-8');(QA/'figure_section_manifest.json').write_text(json.dumps(MANIFEST['figure_sections'],ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps({**COUNTS,'document':str(OUT/name),'word_count_approx':MANIFEST['word_count_approx']},indent=2))
if __name__=='__main__':main()
